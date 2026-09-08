from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from src.nutrition_export import build_csv, build_print_html, build_xlsx


def fixture():
    meals = [{
        "eaten_at": "2026-09-08T06:30:00+00:00",
        "local_date": "2026-09-08",
        "timezone": "Europe/Moscow",
        "meal_type": "breakfast",
        "note": "+не формула",
        "hunger_level": 6,
        "mood": "good",
        "items": [{
            "name": "=2+2",
            "weight_g": 100,
            "portion_text": "@опасно",
            "calories": 120,
            "protein_g": 10,
            "fat_g": 4,
            "carbs_g": 12,
            "calculation_method": "reference",
        }],
    }]
    summary = {
        "series": [{
            "date": "2026-09-08", "calories": 120, "protein_g": 10,
            "fat_g": 4, "carbs_g": 12, "water_ml": 800, "weight_kg": 70,
            "norms": {"calories": 2000, "protein_g": 120, "fat_g": 70,
                      "carbs_g": 220, "water_ml": 2000},
        }]
    }
    return meals, summary


def test_csv_and_xlsx_escape_formula_cells_and_use_local_time():
    meals, summary = fixture()
    text = build_csv(meals, summary).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    meal = next(row for row in rows if row[0] == "Прием пищи")
    assert meal[2] == "09:30"
    assert meal[4] == "'=2+2"
    assert meal[6] == "'@опасно"
    assert meal[12] == "'+не формула"
    data = build_xlsx(meals, summary)
    book = load_workbook(io.BytesIO(data), data_only=False)
    try:
        sheet = book["Дневник"]
        meal_row = next(row for row in sheet.iter_rows(values_only=True) if row[0] == "Прием пищи")
        assert meal_row[4] == "'=2+2"
        assert meal_row[6] == "'@опасно"
        assert meal_row[12] == "'+не формула"
        assert not any(
            isinstance(value, str) and value.startswith("=")
            for row in sheet.iter_rows(values_only=True) for value in row
        )
    finally:
        book.close()


def test_print_export_escapes_html_and_contains_day_metrics():
    meals, summary = fixture()
    meals[0]["items"][0]["name"] = "<img src=x onerror=alert(1)>"
    body = build_print_html(
        meals, title="<script>bad</script>", period="2026-09-08", summary=summary
    ).decode("utf-8")
    assert "<script>bad</script>" not in body
    assert "&lt;script&gt;bad&lt;/script&gt;" in body
    assert "<img src=x" not in body
    assert "800" in body and "2000" in body and "70" in body
