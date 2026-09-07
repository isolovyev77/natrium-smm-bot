import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from src.nutrition_plan import (
    HEADERS,
    PlanValidationError,
    build_template,
    parse_plan,
)


def workbook_bytes(rows=(), headers=HEADERS, extra_sheet=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "План"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if extra_sheet:
        workbook.create_sheet(extra_sheet)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_parse_complete_plan_and_sort_dates():
    payload = workbook_bytes([
        [date(2026, 2, 1), 2100, 125, 72, 265, None],
        ["01.01.2026", "2000.5", "120,5", 70, 250, 2000],
    ])
    rows = parse_plan(payload)
    assert [row.effective_from for row in rows] == [date(2026, 1, 1), date(2026, 2, 1)]
    assert rows[0].calories == Decimal("2000.5")
    assert rows[0].protein_g == Decimal("120.5")
    assert rows[0].water_ml == 2000
    assert rows[1].water_ml is None
    assert rows[0].source_row == 3


def test_parse_path_and_binary_stream(tmp_path):
    payload = workbook_bytes([[date(2026, 1, 1), 2000, 120, 70, 250, 2000]])
    path = tmp_path / "plan.xlsx"
    path.write_bytes(payload)
    assert parse_plan(path) == parse_plan(io.BytesIO(payload))


@pytest.mark.parametrize(
    "row, message",
    [
        (["31.02.2026", 2000, 120, 70, 250, 2000], "Строка 2"),
        ([date(2026, 1, 1), None, 120, 70, 250, 2000], "не заполнено"),
        ([date(2026, 1, 1), 2000, -1, 70, 250, 2000], "не может быть отрицательным"),
        ([date(2026, 1, 1), 2000, 120, "NaN", 250, 2000], "конечным числом"),
        ([date(2026, 1, 1), 2000, 120, 70, 250, 2000.5], "целым числом"),
    ],
)
def test_invalid_rows_are_rejected_with_row_number(row, message):
    with pytest.raises(PlanValidationError, match=message) as error:
        parse_plan(workbook_bytes([row]))
    assert error.value.row_number == 2


def test_duplicate_dates_reject_whole_import():
    payload = workbook_bytes([
        [date(2026, 1, 1), 2000, 120, 70, 250, 2000],
        ["2026-01-01", 2100, 125, 72, 265, 2100],
    ])
    with pytest.raises(PlanValidationError, match="Строка 3: дата начала повторяется"):
        parse_plan(payload)


def test_formula_is_rejected_instead_of_evaluated():
    payload = workbook_bytes([[date(2026, 1, 1), "=1000+1000", 120, 70, 250, 2000]])
    with pytest.raises(PlanValidationError, match="формулы не поддерживаются"):
        parse_plan(payload)


def test_wrong_headers_empty_plan_and_wrong_extension_are_rejected(tmp_path):
    with pytest.raises(PlanValidationError, match="заголовки"):
        parse_plan(workbook_bytes([], headers=("Date", *HEADERS[1:])))
    with pytest.raises(PlanValidationError, match="нет строк"):
        parse_plan(workbook_bytes())
    wrong = tmp_path / "plan.xlsm"
    wrong.write_bytes(workbook_bytes([[date(2026, 1, 1), 1, 1, 1, 1, 1]]))
    with pytest.raises(PlanValidationError, match="только файл .xlsx"):
        parse_plan(wrong)


def test_limits_and_invalid_container_are_rejected():
    with pytest.raises(PlanValidationError, match="2 МБ"):
        parse_plan(b"x" * (2 * 1024 * 1024 + 1))
    with pytest.raises(PlanValidationError, match="корректной книгой"):
        parse_plan(b"not an xlsx")

    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * (10 * 1024 * 1024 + 1))
    with pytest.raises(PlanValidationError, match="распакованный файл"):
        parse_plan(bomb.getvalue())


@pytest.mark.parametrize(
    "row",
    [
        [date(2026, 1, 1), "1e10000", 120, 70, 250, 2000],
        [date(2026, 1, 1), 2000, 2001, 70, 250, 2000],
        [date(2026, 1, 1), 2000, 120, 70, 250, 20001],
    ],
)
def test_large_finite_values_are_rejected_before_preview(row):
    with pytest.raises(PlanValidationError, match="Строка 2: .*технический предел"):
        parse_plan(workbook_bytes([row]))


def test_inflated_worksheet_dimensions_are_rejected_without_iteration():
    payload = workbook_bytes([[date(2026, 1, 1), 1, 1, 1, 1, 1]])
    source = zipfile.ZipFile(io.BytesIO(payload))
    inflated = io.BytesIO()
    with zipfile.ZipFile(inflated, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:F2"', b'ref="A1:XFD1048576"')
            target.writestr(item, content)
    source.close()
    with pytest.raises(PlanValidationError, match="размер листа"):
        parse_plan(inflated.getvalue())


def test_external_links_and_too_many_sheets_are_rejected():
    payload = workbook_bytes([[date(2026, 1, 1), 1, 1, 1, 1, 1]])
    source = zipfile.ZipFile(io.BytesIO(payload))
    linked = io.BytesIO()
    with zipfile.ZipFile(linked, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    source.close()
    with pytest.raises(PlanValidationError, match="внешние ссылки"):
        parse_plan(linked.getvalue())

    workbook = Workbook()
    workbook.active.title = "План"
    workbook.active.append(HEADERS)
    workbook.active.append([date(2026, 1, 1), 1, 1, 1, 1, 1])
    for index in range(5):
        workbook.create_sheet(f"Лист {index}")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(PlanValidationError, match="не больше 5 листов"):
        parse_plan(output.getvalue())


def test_template_is_editable_styled_and_imports_only_plan():
    payload = build_template().getvalue()
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    assert workbook.sheetnames == ["План", "Пример", "Инструкция"]
    assert tuple(cell.value for cell in workbook["План"][1]) == HEADERS
    assert workbook["План"].freeze_panes == "A2"
    assert workbook["План"]["A1"].font.name == "Arial"
    assert workbook["Пример"]["A4"].value.date() == date(2030, 1, 1)
    assert workbook["Пример"]["A1"].value == "Учебный пример. Замените числа нормами вашего клиента."
    assert not any(
        cell.data_type == "f"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    workbook["План"].append([date(2026, 1, 1), 2000, 120, 70, 250, 2000])
    updated = io.BytesIO()
    workbook.save(updated)
    workbook.close()
    rows = parse_plan(updated.getvalue())
    assert len(rows) == 1
    assert rows[0].effective_from == date(2026, 1, 1)
