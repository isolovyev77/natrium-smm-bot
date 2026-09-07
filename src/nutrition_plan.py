"""Safe import and template generation for client nutrition norm plans."""

from __future__ import annotations

import io
import math
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import BinaryIO


MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_UNPACKED_BYTES = 10 * 1024 * 1024
MAX_ZIP_ENTRIES = 100
MAX_SHEETS = 5
MAX_PLAN_ROWS = 366
MAX_CALORIES = Decimal("20000")
MAX_MACRO_GRAMS = Decimal("2000")
MAX_WATER_ML = Decimal("20000")
PLAN_SHEET = "План"
HEADERS = (
    "Дата начала",
    "Калории (ккал/день)",
    "Белки (г/день)",
    "Жиры (г/день)",
    "Углеводы (г/день)",
    "Вода (мл/день)",
)


class NutritionPlanError(Exception):
    """Base error for nutrition plan workbooks."""


class PlanValidationError(NutritionPlanError):
    """A workbook or a particular plan row is invalid."""

    def __init__(self, message: str, row_number: int | None = None):
        self.row_number = row_number
        if row_number is not None:
            message = f"Строка {row_number}: {message}"
        super().__init__(message)


@dataclass(frozen=True)
class PlanRow:
    effective_from: date
    calories: Decimal
    protein_g: Decimal
    fat_g: Decimal
    carbs_g: Decimal
    water_ml: int | None
    source_row: int


def _read_source(source: bytes | bytearray | BinaryIO | Path | str) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        payload = bytes(source)
    elif isinstance(source, (str, Path)):
        path = Path(source)
        if path.suffix.casefold() != ".xlsx":
            raise PlanValidationError("поддерживается только файл .xlsx")
        try:
            size = path.stat().st_size
        except OSError as error:
            raise PlanValidationError("не удалось прочитать файл") from error
        if size > MAX_FILE_BYTES:
            raise PlanValidationError("файл больше допустимых 2 МБ")
        try:
            payload = path.read_bytes()
        except OSError as error:
            raise PlanValidationError("не удалось прочитать файл") from error
    elif hasattr(source, "read"):
        try:
            payload = source.read(MAX_FILE_BYTES + 1)
        except (OSError, TypeError) as error:
            raise PlanValidationError("не удалось прочитать файл") from error
        if not isinstance(payload, (bytes, bytearray)):
            raise PlanValidationError("ожидались двоичные данные файла .xlsx")
        payload = bytes(payload)
    else:
        raise PlanValidationError("ожидался файл .xlsx или его двоичные данные")
    if len(payload) > MAX_FILE_BYTES:
        raise PlanValidationError("файл больше допустимых 2 МБ")
    if not payload:
        raise PlanValidationError("файл пуст")
    return payload


def _check_xlsx_container(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise PlanValidationError("в книге слишком много внутренних файлов")
            total_size = sum(entry.file_size for entry in entries)
            if total_size > MAX_UNPACKED_BYTES:
                raise PlanValidationError("распакованный файл превышает лимит 10 МБ")
            for entry in entries:
                normalized = entry.filename.replace("\\", "/")
                if normalized.startswith("/") or ".." in normalized.split("/"):
                    raise PlanValidationError("некорректная структура файла .xlsx")
                lowered = normalized.casefold()
                if "vbaproject.bin" in lowered:
                    raise PlanValidationError("книги с макросами не поддерживаются")
                if lowered.startswith("xl/externallinks/"):
                    raise PlanValidationError("внешние ссылки в книге не поддерживаются")
                if entry.flag_bits & 0x1:
                    raise PlanValidationError("зашифрованные книги не поддерживаются")
                if lowered.endswith(".rels") and entry.file_size <= 512 * 1024:
                    with archive.open(entry) as stream:
                        rels = stream.read(512 * 1024 + 1).lower()
                    if b'targetmode="external"' in rels or b"targetmode='external'" in rels:
                        raise PlanValidationError("внешние ссылки в книге не поддерживаются")
    except PlanValidationError:
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError) as error:
        raise PlanValidationError("файл не является корректной книгой .xlsx") from error


def _parse_date(value: object, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    raise PlanValidationError("укажите дату в формате ДД.ММ.ГГГГ", row_number)


def _parse_decimal(
    value: object, label: str, row_number: int, maximum: Decimal
) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise PlanValidationError(f"не заполнено поле «{label}»", row_number)
    if isinstance(value, bool):
        raise PlanValidationError(f"поле «{label}» должно быть числом", row_number)
    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError, AttributeError):
        raise PlanValidationError(f"поле «{label}» должно быть числом", row_number)
    if not number.is_finite():
        raise PlanValidationError(f"поле «{label}» должно быть конечным числом", row_number)
    if number < 0:
        raise PlanValidationError(f"поле «{label}» не может быть отрицательным", row_number)
    if number > maximum:
        raise PlanValidationError(
            f"поле «{label}» превышает технический предел {maximum}", row_number
        )
    return number


def parse_plan(source: bytes | bytearray | BinaryIO | Path | str) -> list[PlanRow]:
    """Validate the whole workbook and return plan rows without side effects."""
    payload = _read_source(source)
    _check_xlsx_container(payload)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
    except ImportError as error:
        raise NutritionPlanError("для импорта .xlsx не установлен openpyxl") from error
    except Exception as error:
        raise PlanValidationError("не удалось открыть книгу .xlsx") from error

    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise PlanValidationError(f"в книге должно быть не больше {MAX_SHEETS} листов")
        if PLAN_SHEET not in workbook.sheetnames:
            raise PlanValidationError(f"не найден лист «{PLAN_SHEET}»")
        sheet = workbook[PLAN_SHEET]
        if sheet.max_row > MAX_PLAN_ROWS + 1 or sheet.max_column > len(HEADERS):
            raise PlanValidationError(
                f"размер листа «План» превышает {MAX_PLAN_ROWS} строк и {len(HEADERS)} колонок"
            )
        iterator = sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, MAX_PLAN_ROWS + 1),
            min_col=1,
            max_col=len(HEADERS),
        )
        try:
            header_cells = next(iterator)
        except StopIteration:
            raise PlanValidationError("лист «План» пуст")
        actual_headers = tuple(cell.value for cell in header_cells[: len(HEADERS)])
        if actual_headers != HEADERS:
            raise PlanValidationError("заголовки листа «План» не соответствуют шаблону", 1)
        rows: list[PlanRow] = []
        seen_dates: set[date] = set()
        for row_number, cells in enumerate(iterator, start=2):
            if any(cell.data_type == "f" for cell in cells):
                raise PlanValidationError("формулы не поддерживаются", row_number)
            values = [cell.value for cell in cells]
            if not any(value not in (None, "") for value in values):
                continue
            if len(rows) >= MAX_PLAN_ROWS:
                raise PlanValidationError(f"допустимо не больше {MAX_PLAN_ROWS} строк плана")
            values.extend([None] * (len(HEADERS) - len(values)))
            effective_from = _parse_date(values[0], row_number)
            if effective_from in seen_dates:
                raise PlanValidationError("дата начала повторяется", row_number)
            seen_dates.add(effective_from)
            calories = _parse_decimal(values[1], HEADERS[1], row_number, MAX_CALORIES)
            protein = _parse_decimal(values[2], HEADERS[2], row_number, MAX_MACRO_GRAMS)
            fat = _parse_decimal(values[3], HEADERS[3], row_number, MAX_MACRO_GRAMS)
            carbs = _parse_decimal(values[4], HEADERS[4], row_number, MAX_MACRO_GRAMS)
            water = None
            if values[5] not in (None, ""):
                water_value = _parse_decimal(values[5], HEADERS[5], row_number, MAX_WATER_ML)
                if water_value != water_value.to_integral_value():
                    raise PlanValidationError("вода должна быть указана целым числом миллилитров", row_number)
                water = int(water_value)
            rows.append(PlanRow(effective_from, calories, protein, fat, carbs, water, row_number))
        if not rows:
            raise PlanValidationError("на листе «План» нет строк для импорта")
        return sorted(rows, key=lambda row: row.effective_from)
    finally:
        workbook.close()


def build_template() -> io.BytesIO:
    """Build an editable, formula-free workbook matching parse_plan()."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as error:
        raise NutritionPlanError("для создания шаблона не установлен openpyxl") from error

    workbook = Workbook()
    plan = workbook.active
    plan.title = PLAN_SHEET
    example = workbook.create_sheet("Пример")
    instruction = workbook.create_sheet("Инструкция")
    header_fill = PatternFill("solid", fgColor="285A64")
    input_fill = PatternFill("solid", fgColor="EAF4F3")
    thin = Side(style="thin", color="B8C7C9")
    widths = (17, 24, 21, 21, 25, 20)

    def style_table(sheet, header_row: int, last_row: int) -> None:
        for cell in sheet[header_row]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[header_row].height = 36
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        for row in sheet.iter_rows(min_row=header_row + 1, max_row=last_row, max_col=len(HEADERS)):
            for cell in row:
                cell.font = Font(name="Arial", color="1F2933")
                cell.fill = input_fill
                cell.border = Border(bottom=thin)
            row[0].number_format = "DD.MM.YYYY"
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:F{max(header_row, last_row)}"

    plan.append(HEADERS)
    plan.append([None] * len(HEADERS))
    style_table(plan, 1, 50)
    plan.sheet_view.showGridLines = False

    example["A1"] = "Учебный пример. Замените числа нормами вашего клиента."
    example["A1"].font = Font(name="Arial", bold=True, color="7A4B00")
    example["A1"].fill = PatternFill("solid", fgColor="FFF1CC")
    example.merge_cells("A1:F1")
    example["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    example.row_dimensions[1].height = 34
    for column, header in enumerate(HEADERS, start=1):
        example.cell(3, column, header)
    example.append([date(2030, 1, 1), 2000, 120, 70, 250, 2000])
    example.append([date(2030, 2, 1), 2100, 125, 72, 265, 2100])
    style_table(example, 3, 5)
    example.sheet_view.showGridLines = False

    instructions = (
        ("Как заполнить", "Вносите данные только на листе «План», начиная со строки 2."),
        ("Период действия", "Каждая строка действует с указанной даты до даты следующей строки."),
        ("Обязательные поля", "Дата начала, калории, белки, жиры и углеводы. Вода может быть пустой."),
        ("Единицы", "Калории: ккал/день; белки, жиры и углеводы: г/день; вода: мл/день."),
        ("Ограничения", "Не добавляйте формулы, ссылки, макросы и другие колонки. Не меняйте заголовки."),
        ("Пример", "Лист «Пример» не импортируется. Его числа учебные и не являются медицинской рекомендацией."),
    )
    instruction.append(("Раздел", "Пояснение"))
    for row in instructions:
        instruction.append(row)
    instruction.column_dimensions["A"].width = 23
    instruction.column_dimensions["B"].width = 88
    instruction.freeze_panes = "A2"
    instruction.sheet_view.showGridLines = False
    for cell in instruction[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row in instruction.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", color="1F2933")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        instruction.row_dimensions[row[0].row].height = 36

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
