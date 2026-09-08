"""Safe diary exports without Telegram identifiers or executable cells."""

from __future__ import annotations

import csv
import html
import io
from datetime import datetime
from typing import Any, Iterable
from zoneinfo import ZoneInfo


HEADERS = (
    "Тип строки",
    "Дата",
    "Время",
    "Прием пищи",
    "Продукт",
    "Масса, г",
    "Порция",
    "Калории, ккал",
    "Белки, г",
    "Жиры, г",
    "Углеводы, г",
    "Метод расчета",
    "Заметка",
    "Голод, 1-10",
    "Настроение",
    "Вода, мл",
    "Вес, кг",
    "Норма калорий, ккал/день",
    "Норма белков, г/день",
    "Норма жиров, г/день",
    "Норма углеводов, г/день",
    "Норма воды, мл/день",
)

MEAL_TYPES = {
    "breakfast": "Завтрак", "lunch": "Обед", "dinner": "Ужин",
    "snack": "Перекус", "завтрак": "Завтрак", "обед": "Обед",
    "ужин": "Ужин", "перекус": "Перекус",
}
METHODS = {
    "reference": "По справочнику", "manual": "Вручную",
    "ai": "Приблизительная оценка ИИ", "legacy": "Старая запись",
}


def _safe_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.replace("\x00", "")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def meal_rows(meals: Iterable[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for meal in meals:
        eaten_at = str(meal.get("eaten_at") or "")
        time_value = ""
        try:
            value = datetime.fromisoformat(eaten_at.replace("Z", "+00:00"))
            time_value = value.astimezone(ZoneInfo(meal.get("timezone") or "Europe/Moscow")).strftime("%H:%M")
        except (ValueError, TypeError):
            pass
        items = meal.get("items") or []
        for item in items:
            rows.append(
                [
                    "Прием пищи",
                    meal.get("local_date", ""),
                    time_value,
                    MEAL_TYPES.get(str(meal.get("meal_type", "")).lower(), meal.get("meal_type", "")),
                    item.get("name", ""),
                    item.get("weight_g"),
                    item.get("portion_text", ""),
                    item.get("calories"),
                    item.get("protein_g"),
                    item.get("fat_g"),
                    item.get("carbs_g"),
                    METHODS.get(item.get("calculation_method", ""), item.get("calculation_method", "")),
                    meal.get("note", ""),
                    meal.get("hunger_level"),
                    meal.get("mood"),
                    None, None, None, None, None, None, None,
                ]
            )
    return [[_safe_cell(value) for value in row] for row in rows]


def report_rows(
    meals: Iterable[dict[str, Any]], summary: dict[str, Any] | None = None
) -> list[list[Any]]:
    rows = meal_rows(meals)
    for day in (summary or {}).get("series", []):
        norms = day.get("norms") or {}
        rows.append([
            "Итог дня", day.get("date"), "", "", "", None, "",
            day.get("calories"), day.get("protein_g"), day.get("fat_g"),
            day.get("carbs_g"), "", "", None, "", day.get("water_ml"),
            day.get("weight_kg"), norms.get("calories"), norms.get("protein_g"),
            norms.get("fat_g"), norms.get("carbs_g"), norms.get("water_ml"),
        ])
    return [[_safe_cell(value) for value in row] for row in sorted(rows, key=lambda row: (str(row[1]), str(row[2]), str(row[0])))]


def build_csv(
    meals: Iterable[dict[str, Any]], summary: dict[str, Any] | None = None
) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(HEADERS)
    writer.writerows(report_rows(meals, summary))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def build_xlsx(
    meals: Iterable[dict[str, Any]], summary: dict[str, Any] | None = None
) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Для экспорта XLSX не установлен openpyxl") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Дневник"
    sheet.append(HEADERS)
    for row in report_rows(meals, summary):
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285A64")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (14, 13, 9, 18, 32, 12, 20, 16, 12, 10, 15, 24, 36, 14, 18, 12, 12, 20, 19, 17, 23, 19)
    from openpyxl.utils import get_column_letter
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_print_html(
    meals: Iterable[dict[str, Any]],
    *,
    title: str,
    period: str,
    summary: dict[str, Any] | None = None,
) -> bytes:
    def escaped(value: Any) -> str:
        return html.escape("" if value is None else str(value), quote=True)

    body_rows = []
    for row in report_rows(meals, summary):
        body_rows.append("<tr>" + "".join(f"<td>{escaped(value)}</td>" for value in row) + "</tr>")
    document = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><title>{escaped(title)}</title>
<style>body{{font:14px Arial,sans-serif;color:#1f2933}}table{{border-collapse:collapse;width:100%}}
th,td{{border:1px solid #cbd5e1;padding:5px;text-align:left}}th{{background:#e8f2f3}}
@media print{{button{{display:none}}}}</style></head><body>
<button onclick="window.print()">Печать или PDF</button>
<h1>{escaped(title)}</h1><p>{escaped(period)}</p>
<table><thead><tr>{''.join(f'<th>{escaped(value)}</th>' for value in HEADERS)}</tr></thead>
<tbody>{''.join(body_rows)}</tbody></table></body></html>"""
    return document.encode("utf-8")
